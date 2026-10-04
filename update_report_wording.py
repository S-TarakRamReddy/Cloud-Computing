import os
from docx import Document

base_dir = os.path.dirname(__file__)
report_path = os.path.join(base_dir, "Cloud_Based_Credit_Card_Fraud_Detection_Report.docx")

doc = Document(report_path)

# Find Chapter 11 and add the text after it.
# Actually, appending it to the end of Chapter 11. 
# It's easier to just find the paragraph that starts with "Actual Confusion Matrix" and insert before it,
# or just append it at the end of the document. The user says "Update the report wording to say:"
# I will just add a new paragraph right before Chapter 12.

for i, p in enumerate(doc.paragraphs):
    if p.text.startswith("Chapter 12"):
        p.insert_paragraph_before("Probability calibration improved the reliability of the model's probability estimates, as reflected by the lower Brier score, while preserving the classification metrics at the selected operating threshold.")
        break
else:
    doc.add_paragraph("Probability calibration improved the reliability of the model's probability estimates, as reflected by the lower Brier score, while preserving the classification metrics at the selected operating threshold.")

doc.save(report_path)
print("Report updated.")
